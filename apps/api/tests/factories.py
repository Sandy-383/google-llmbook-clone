"""Builds real test documents in-process, so tests need no binary fixtures."""

import asyncio
import hashlib
import io
import json
import math
import re
import uuid
from collections.abc import AsyncIterator
from typing import cast

import docx
from PIL import Image, ImageDraw, ImageFont

from app.db.models import EMBED_DIM
from app.llm.ollama import ChatChunk, ChatStats, OllamaClient


def parse_sse(body: str) -> list[tuple[str, dict[str, object]]]:
    events = []
    for block in body.strip().split("\n\n"):
        event, data = "", {}
        for line in block.splitlines():
            if line.startswith("event: "):
                event = line.removeprefix("event: ")
            elif line.startswith("data: "):
                data = json.loads(line.removeprefix("data: "))
        events.append((event, data))
    return events


def make_pdf(pages: list[list[str]]) -> bytes:
    """A minimal valid PDF with a real text layer: one Helvetica line per string."""
    page_ids = [4 + 2 * i for i in range(len(pages))]
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for page_id, lines in zip(page_ids, pages, strict=True):
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {page_id + 1} 0 R >>".encode()
        )
        ops = ["BT", "/F1 11 Tf", "14 TL", "72 740 Td"]
        for line in lines:
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            ops.append(f"({escaped}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    ).encode()
    return bytes(out)


def make_scanned_pdf(text: str) -> bytes:
    """An image-only page, like a scan: no text layer at all, so it needs OCR."""
    image = Image.new("RGB", (1700, 2200), "white")
    draw = ImageDraw.Draw(image)
    draw.multiline_text(
        (120, 160), text, fill="black", font=ImageFont.load_default(size=56), spacing=24
    )
    buffer = io.BytesIO()
    image.save(buffer, format="PDF", resolution=200.0)
    return buffer.getvalue()


def make_docx() -> bytes:
    document = docx.Document()
    document.add_heading("Employee Handbook", level=0)
    document.add_heading("Travel", level=1)
    document.add_paragraph("Economy class is required for flights under 6 hours.")
    document.add_heading("Leave", level=1)
    document.add_paragraph("Full-time employees receive 27 days of paid annual leave.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Leave type", "Days"
    table.cell(1, 0).text, table.cell(1, 1).text = "Parental", "112"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


HANDBOOK_PAGES = [
    [
        "Halvorsen Instruments Handbook",
        "Travel: economy class is required for flights under 6 hours.",
        "Business class is allowed for longer flights with director approval.",
        "Page 1",
    ],
    [
        "Halvorsen Instruments Handbook",
        "Annual leave: full-time employees receive 27 days of paid annual leave.",
        "Up to 5 unused days may be carried over into the next year.",
        "Page 2",
    ],
    [
        "Halvorsen Instruments Handbook",
        "Security: all laptops must use full-disk encryption.",
        "Laptops lock automatically after 5 minutes of inactivity.",
        "Page 3",
    ],
    [
        "Halvorsen Instruments Handbook",
        "Expenses: a receipt is required for any single expense over 25 euros.",
        "Claims must be submitted within 30 days.",
        "Page 4",
    ],
]


def fake_vector(text: str) -> list[float]:
    """Deterministic stand-in for an embedding: a hashed bag of words, normalised. Texts that
    share words get a real cosine similarity, so retrieval can be tested without a model."""
    vector = [0.0] * EMBED_DIM
    for word in re.findall(r"[a-z0-9]+", text.lower()):
        vector[int(hashlib.sha256(word.encode()).hexdigest(), 16) % EMBED_DIM] += 1.0
    norm = math.sqrt(sum(x * x for x in vector))
    if norm == 0:  # no words: avoid a zero vector, whose cosine distance is undefined
        vector[0], norm = 1.0, 1.0
    return [x / norm for x in vector]


async def fake_embed(texts: list[str]) -> list[list[float]]:
    return [fake_vector(t) for t in texts]


class FakeOllama:
    """Stands in for the model server in chat tests. Streams `reply` word by word."""

    def __init__(
        self,
        reply: str = "Employees get 27 days of paid annual leave [1].",
        rewrite: str = "annual leave carry over days",
        fail: Exception | None = None,
        delay: float = 0.0,
    ) -> None:
        self.reply, self.rewrite, self.fail, self.delay = reply, rewrite, fail, delay
        self.answer_requests: list[list[dict[str, str]]] = []
        self.rewrite_requests: list[list[dict[str, str]]] = []
        self.embedded: list[str] = []
        # Non-streamed calls (query rewrites, guides, studio): scripted replies, used in order.
        self.chat_replies: list[str] = []
        self.chat_formats: list[object] = []

    async def missing_models(self, required: list[str]) -> list[str]:
        return []

    async def embed(
        self, texts: list[str], *, model: str, keep_alive: str, cpu_only: bool = False
    ) -> list[list[float]]:
        self.embedded += texts
        return [fake_vector(t) for t in texts]

    async def chat(self, messages: list[dict[str, str]], format: object = None, **_: object) -> str:
        self.rewrite_requests.append(messages)
        self.chat_formats.append(format)
        if self.fail:
            raise self.fail
        return self.chat_replies.pop(0) if self.chat_replies else self.rewrite

    async def chat_stream(
        self, messages: list[dict[str, str]], **_: object
    ) -> AsyncIterator[ChatChunk]:
        self.answer_requests.append(messages)
        if self.fail:
            raise self.fail
        for word in re.findall(r"\S+\s*", self.reply):
            if self.delay:
                await asyncio.sleep(self.delay)
            yield ChatChunk(text=word)
        yield ChatChunk(stats=ChatStats(400, 12, 0, 20, 300))


class RecordingQueue:
    """Stands in for Celery in API tests: records what would have been enqueued."""

    def __init__(self) -> None:
        self.ingested: list[uuid.UUID] = []
        self.guides: list[uuid.UUID] = []
        self.studios: list[uuid.UUID] = []

    async def enqueue_ingest(self, source_id: uuid.UUID) -> None:
        self.ingested.append(source_id)

    async def enqueue_guide(self, source_id: uuid.UUID) -> None:
        self.guides.append(source_id)

    async def enqueue_studio(self, note_id: uuid.UUID) -> None:
        self.studios.append(note_id)


def as_client(fake: FakeOllama) -> OllamaClient:
    """FakeOllama implements the parts of OllamaClient that the code under test uses."""
    return cast(OllamaClient, fake)
